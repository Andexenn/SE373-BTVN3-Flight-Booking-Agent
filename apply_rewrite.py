import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_paragraph_content(p, runs_data, align=None, space_after=4, line_spacing=1.15):
    """
    Clears text in paragraph and adds runs with rich formatting.
    runs_data: list of tuples (text, is_bold, is_italic, font_size)
    """
    p.text = ""
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

    for item in runs_data:
        text = item[0]
        is_bold = item[1] if len(item) > 1 else False
        is_italic = item[2] if len(item) > 2 else False
        size_pt = item[3] if len(item) > 3 else 12.0
        
        r = p.add_run(text)
        r.bold = is_bold
        r.italic = is_italic
        if size_pt:
            r.font.size = Pt(size_pt)

def update_document():
    docx_path = "BTTH_Buoi3_SE373_PhanHuynhHung_24520623.docx"
    doc = docx.Document(docx_path)

    # -------------------------------------------------------------
    # COVER / HEADER UPDATES
    # -------------------------------------------------------------
    set_paragraph_content(doc.paragraphs[1], [("Trường Đại học Công nghệ Thông tin", True, False, 15.0)], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    set_paragraph_content(doc.paragraphs[2], [("Đại học Quốc gia Tp.Hồ Chí Minh", True, False, 14.0)], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    set_paragraph_content(doc.paragraphs[3], [("Toàn diện – Sáng tạo - Phụng sự", False, True, 12.0)], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # Paragraph 4 is drawing (logo) - DO NOT TOUCH

    set_paragraph_content(doc.paragraphs[5], [("BÁO CÁO BÀI TẬP THỰC HÀNH BUỔI 3", True, False, 18.0)], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    set_paragraph_content(doc.paragraphs[7], [("Môn học: ", True, False, 13.0), ("Kỹ thuật xây dựng hệ thống Agentic AI (SE373)", False, False, 13.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[8], [("Bài tập: ", True, False, 13.0), ("Dựng Agent đặt vé máy bay bằng LangChain & Harness an toàn", False, False, 13.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[9], [("Họ và tên sinh viên: ", True, False, 13.0), ("Phan Huỳnh Hưng", False, False, 13.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[10], [("Mã số sinh viên: ", True, False, 13.0), ("24520623", False, False, 13.0)], space_after=16)

    # Clear unnecessary excessive blank paragraphs (P11 to P24)
    for i in range(11, 25):
        doc.paragraphs[i].text = ""
        doc.paragraphs[i].paragraph_format.space_after = Pt(0)
        doc.paragraphs[i].paragraph_format.line_spacing = 1.0

    # -------------------------------------------------------------
    # SECTION 1: THIẾT KẾ CÁC LỚP HARNESS BẢO VỆ
    # -------------------------------------------------------------
    set_paragraph_content(doc.paragraphs[25], [("1) THIẾT KẾ CÁC LỚP HARNESS BẢO VỆ (HARNESS SAFETY LAYERS)", True, False, 13.5)], space_after=8)

    set_paragraph_content(doc.paragraphs[26], [
        ("• Nguyên lý cốt lõi: ", True, False, 12.0),
        ("Hệ thống tuân thủ nghiêm ngặt công thức nền tảng của kỹ nghệ Agent: ", False, False, 12.0),
        ("Agent = Model + Harness", True, False, 12.0),
        (". Trong đó, Model đảm nhiệm năng lực suy luận (reasoning) và đề xuất công cụ, còn Harness là lớp vỏ bọc mã nguồn cố định (deterministic code wrapper) độc lập, chịu trách nhiệm thiết lập kỷ luật dữ liệu, kiểm soát thẩm quyền giao dịch và nghiệm thu kết quả.", False, False, 12.0)
    ], space_after=6)

    doc.paragraphs[27].text = ""
    doc.paragraphs[27].paragraph_format.space_after = Pt(0)

    set_paragraph_content(doc.paragraphs[28], [
        ("• Kiến trúc 4 lớp bảo vệ: ", True, False, 12.0),
        ("Hệ thống đã module hóa toàn bộ cơ chế bảo vệ trong package harness/ bao gồm 4 tầng phân lập rõ ràng:", False, False, 12.0)
    ], space_after=6)

    # Layer 1
    set_paragraph_content(doc.paragraphs[29], [("Lớp 1: Ràng buộc là dữ liệu (harness/constraints.py - Constraints as Data)", True, False, 12.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[30], [
        ("  - Bản chất: ", True, False, 12.0),
        ("Chuyển hóa toàn bộ yêu cầu đặt vé của người dùng từ văn bản tự do (free-form prompt) thành mô hình dữ liệu Pydantic BookingConstraint có cấu trúc định kiểu chặt chẽ (strict typing), triệt tiêu hoàn toàn hiện tượng trôi mục tiêu (goal drift).", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[31], [
        ("  - Cấu trúc: ", True, False, 12.0),
        ("Quản lý định lượng các trường: origin, dest (IATA 3 ký tự), date (YYYY-MM-DD), max_price (ngân sách trần), depart_before (giờ cất cánh), passenger_name, baggage_kg (hành lý) và require_refundable.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[32], [
        ("  - Tác dụng: ", True, False, 12.0),
        ("Cung cấp phương thức to_natural_language() định hình ngữ cảnh chuẩn xác cho Model, đồng thời làm hệ quy chiếu chân lý (Ground Truth) bất biến phục vụ kiểm định nghiệm thu tự động.", False, False, 12.0)
    ], space_after=6)

    # Layer 2
    set_paragraph_content(doc.paragraphs[33], [("Lớp 2: Kiểm quyền trước khi gọi công cụ (harness/permission.py - Permission Guard)", True, False, 12.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[34], [
        ("  - Vị trí can thiệp: ", True, False, 12.0),
        ("Can thiệp chặn trước (Pre-execution Interception) tại cổng dispatcher trước khi bất kỳ công cụ nào được kích hoạt thực thi.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[35], [
        ("  - Cơ chế phân quyền kép (Dual-tier Authorization):", True, False, 12.0)
    ], space_after=2)
    set_paragraph_content(doc.paragraphs[36], [
        ("    + Công cụ chỉ đọc (Read-only): ", True, False, 12.0),
        ("search_flights, check_seat_availability, calculate_baggage_fee chỉ mang tính tra cứu nên được tự động cấp phép (0% độ trễ cho người dùng).", False, False, 12.0)
    ], space_after=2)
    set_paragraph_content(doc.paragraphs[37], [
        ("    + Công cụ nhạy cảm tài chính (Sensitive Financial Action): ", True, False, 12.0),
        ("book_ticket làm biến đổi cơ sở dữ liệu và trừ ngân sách, bắt buộc kích hoạt chốt chặn Human-in-the-loop CLI ([y/n]) để người dùng xác nhận trực tiếp.", False, False, 12.0)
    ], space_after=2)
    set_paragraph_content(doc.paragraphs[38], [
        ("    + Cảnh báo chính sách: ", True, False, 12.0),
        ("Cảnh báo rõ ràng nếu vé không hoàn tiền (refundable == False) và từ chối tức thì nếu tổng chi phí vượt quá trần ngân sách cho phép.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[39], [
        ("  - Tác dụng: ", True, False, 12.0),
        ("Loại bỏ 100% rủi ro phát sinh giao dịch tài chính ngoài ý muốn, đồng thời hỗ trợ tham số auto_approve cho các kịch bản chạy benchmark tự động.", False, False, 12.0)
    ], space_after=6)

    # Layer 3
    set_paragraph_content(doc.paragraphs[40], [("Lớp 3: Tiêu chí hoàn thành kiểm bằng code (harness/completion_sensor.py - Programmatic Sensor)", True, False, 12.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[41], [
        ("  - Bản chất: ", True, False, 12.0),
        ("Áp dụng Cảm biến tính toán (Computational Sensor) kiểm chứng độc lập bằng code logic trên cơ sở dữ liệu backend, hoàn toàn tách biệt và không tin vào phát ngôn chủ quan của LLM (Inferential Sensor).", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[42], [
        ("  - Logic thẩm định chân lý: ", True, False, 12.0),
        ("Hợp lệ ⇔ (Mã vé ∈ MockDB) ∧ (Status == 'confirmed') ∧ (Tên hành khách == BookingConstraint) ∧ (Hành trình, Ngày bay khớp) ∧ (Giờ bay < depart_before) ∧ (Tổng giá vé + hành lý ≤ max_price).", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[43], [
        ("  - Tác dụng: ", True, False, 12.0),
        ("Triệt tiêu hoàn toàn hiện tượng False-Positive Completion (Model tự ảo giác xác nhận đã đặt vé thành công trong khi backend chưa ghi nhận hoặc đặt sai dữ liệu).", False, False, 12.0)
    ], space_after=6)

    # Layer 4
    set_paragraph_content(doc.paragraphs[44], [("Lớp 4: Cơ chế bàn giao con người (harness/handoff.py) & Phát hiện lặp (harness/loop_detector.py)", True, False, 12.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[45], [
        ("  - Bộ phát hiện lặp (LoopDetector): ", True, False, 12.0),
        ("Lưu vết băm (tool_name, hash(args)) trong hàng đợi trượt. Nếu cùng một hành vi bị lặp lại quá ngưỡng K=2 lần mà không tạo tiến triển mới, hệ thống tự động kích hoạt cầu dao ngắt mạch (Circuit Breaker) để tránh lãng phí token và tài nguyên.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[46], [
        ("  - Báo cáo bàn giao (HandoffReport): ", True, False, 12.0),
        ("Khi tiến trình bị dừng do chạm quyền kiểm soát ([y/n] chọn 'n') hoặc phát hiện lặp, hệ thống xuất báo cáo chuẩn Pydantic gồm: Trạng thái, Cột mốc gián đoạn, Lịch sử hành động, Lý do can thiệp và Câu hỏi hành động ngắn gọn giúp con người đưa ra quyết định xử lý chỉ trong 30 giây.", False, False, 12.0)
    ], space_after=8)

    # Paragraph 47 is DRAWING (test_manual screenshot) - DO NOT TOUCH
    set_paragraph_content(doc.paragraphs[48], [("Hình 1: Kết quả kiểm thử độc lập 4 lớp Harness bảo vệ đạt tỷ lệ vượt qua 100% (test_manual.py)", True, False, 10.5)], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    doc.paragraphs[49].text = ""
    doc.paragraphs[50].text = ""

    # -------------------------------------------------------------
    # SECTION 2: CÀI ĐẶT 3 MẪU THIẾT KẾ AGENT
    # -------------------------------------------------------------
    set_paragraph_content(doc.paragraphs[51], [("2) CÀI ĐẶT 3 MẪU THIẾT KẾ AGENT (3 AGENT ARCHITECTURES)", True, False, 13.5)], space_after=8)
    set_paragraph_content(doc.paragraphs[52], [
        ("Cả 3 mẫu thiết kế đều được xây dựng chuẩn hóa trên nền tảng LangChain, hỗ trợ gọi mô hình trực tuyến (OpenAI/Gemini) hoặc bộ suy luận cục bộ (MockReasoningModel) phục vụ đánh giá định lượng:", False, False, 12.0)
    ], space_after=6)

    # Pattern 1: ReAct
    set_paragraph_content(doc.paragraphs[53], [("Mẫu 1: ReAct Agent (agents/react_agent.py - Reasoning + Acting)", True, False, 12.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[54], [
        ("• Nguyên lý: ", True, False, 12.0),
        ("Vận hành theo chu trình phản xạ tuần tự khép kín: Thought → Action → Observation.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[55], [
        ("• Cơ chế thích ứng: ", True, False, 12.0),
        ("Model tự phân tích kết quả quan sát của bước trước để quyết định công cụ kế tiếp. Toàn bộ tool calls đều được bao bọc bởi LoopDetector và PermissionGuard. Khi chuyến bay ưu tiên QH118 cạn kiệt chỗ trống, Agent lập tức suy luận lại và chuyển sang chuyến bay dự phòng VN122 chỉ sau đúng 1 vòng lặp bổ sung.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[56], [
        ("• Đặc điểm: ", True, False, 12.0),
        ("Tốc độ phản xạ tức thì, chi phí suy luận ban đầu thấp (0s trễ lập kế hoạch), đặc biệt phù hợp cho các bài toán đặt vé đơn tuyến có tính biến động cao.", False, False, 12.0)
    ], space_after=6)

    # Pattern 2: Plan-then-Execute
    set_paragraph_content(doc.paragraphs[57], [("Mẫu 2: Plan-then-Execute Agent (agents/plan_execute_agent.py - Lập kế hoạch & Thực thi)", True, False, 12.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[58], [
        ("• Nguyên lý: ", True, False, 12.0),
        ("Phân định rành mạch giữa tư duy chiến lược vĩ mô và thực thi vi mô qua 3 thành phần AI độc lập:", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[59], [
        ("  + LLM Planner: ", True, False, 12.0),
        ("Phân tích ràng buộc và sinh trước bản kế hoạch tổng thể dạng JSON có cấu trúc Pydantic ExecutionPlan gồm 4 bước: search_flights → check_seat_availability → calculate_baggage_fee → book_ticket.", False, False, 12.0)
    ], space_after=2)
    set_paragraph_content(doc.paragraphs[60], [
        ("  + LLM Step Executor: ", True, False, 12.0),
        ("Tiếp nhận từng bước theo kế hoạch, tham chiếu kết quả các bước trước để truyền tham số chính xác nhất cho từng công cụ.", False, False, 12.0)
    ], space_after=2)
    set_paragraph_content(doc.paragraphs[61], [
        ("  + LLM Dynamic Re-planner: ", True, False, 12.0),
        ("Khi phát hiện chuyến bay hết ghế (QH118 hết vé), Executor tạm dừng và chuyển toàn bộ trace lỗi cho Re-planner. Re-planner tái lập ngay kế hoạch phụ phục hồi cục bộ (chuyển sang VN122) mà không làm sụp đổ toàn bộ luồng quy trình.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[62], [
        ("• Đặc điểm: ", True, False, 12.0),
        ("Tính minh bạch và khả năng kiểm toán (Auditability) cực cao do toàn bộ lộ trình được định hình tường minh trước khi chạy.", False, False, 12.0)
    ], space_after=6)

    # Pattern 3: Hybrid
    set_paragraph_content(doc.paragraphs[63], [("Mẫu 3: Mẫu Lai - Hybrid Agent (agents/hybrid_agent.py - Phân tầng vĩ mô & vi mô)", True, False, 12.0)], space_after=4)
    set_paragraph_content(doc.paragraphs[64], [
        ("• Nguyên lý: ", True, False, 12.0),
        ("Áp dụng kiến trúc phân tầng 2 cấp (Two-tier Hierarchical Architecture) kết hợp điều phối vĩ mô với các Micro ReAct Sub-Agents:", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[65], [
        ("  + Tầng điều phối vĩ mô (Milestone Coordinator): ", True, False, 12.0),
        ("Phân chia quy trình đặt vé thành 3 cột mốc chiến lược độc lập: Cột mốc 1: Khám phá chuyến bay (Discovery); Cột mốc 2: Thẩm định chỗ ngồi & cước hành lý (Validation & Pricing); Cột mốc 3: Kiểm quyền & Chốt vé (Booking & Settlement).", False, False, 12.0)
    ], space_after=2)
    set_paragraph_content(doc.paragraphs[66], [
        ("  + Tầng thực thi vi mô (Scoped Micro ReAct Sub-Agents): ", True, False, 12.0),
        ("Tại mỗi cột mốc, một ReAct Sub-Agent độc lập đảm nhiệm giải quyết mục tiêu với không gian ngữ cảnh (Prompt Context) cô lập hoàn toàn. Khi QH118 hết vé ở Cột mốc 2, Sub-Agent 2 tự phục hồi sang VN122 nội bộ trước khi chuyển giao trạng thái hoàn tất (shared_state) sang Cột mốc 3.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[67], [
        ("• Đặc điểm: ", True, False, 12.0),
        ("Khả năng cô lập lỗi tối đa, ngăn ngừa triệt để hiện tượng ô nhiễm ngữ cảnh (Context Pollution) khi hệ thống phải xử lý khối lượng dữ liệu chuyến bay lớn.", False, False, 12.0)
    ], space_after=8)

    # Clear legacy empty paragraphs 68 to 72
    for i in range(68, 73):
        doc.paragraphs[i].text = ""
        doc.paragraphs[i].paragraph_format.space_after = Pt(0)

    # Paragraphs 73, 74 are DRAWINGS (main.py screenshots) - DO NOT TOUCH
    set_paragraph_content(doc.paragraphs[75], [("Hình 2: Cơ chế can thiệp kiểm quyền tương tác người dùng khi chạy main.py", True, False, 10.5)], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # -------------------------------------------------------------
    # SECTION 3: ĐÁNH GIÁ THỰC NGHIỆM & PHÂN TÍCH KẾT QUẢ
    # -------------------------------------------------------------
    set_paragraph_content(doc.paragraphs[76], [("3) ĐÁNH GIÁ THỰC NGHIỆM & PHÂN TÍCH KẾT QUẢ (EMPIRICAL EVALUATION)", True, False, 13.5)], space_after=8)
    set_paragraph_content(doc.paragraphs[77], [
        ("Bảng tổng hợp kết quả đo đạc thực nghiệm từ chương trình benchmark chuẩn hóa (benchmark.py):", False, False, 12.0)
    ], space_after=6)

    # Table 0 is preserved intact between P77 and P78

    # Paragraphs 78 to 80 are empty
    for i in range(78, 81):
        doc.paragraphs[i].text = ""
        doc.paragraphs[i].paragraph_format.space_after = Pt(0)

    # Paragraphs 81, 82, 83 are DRAWINGS (benchmark screenshots) - DO NOT TOUCH
    set_paragraph_content(doc.paragraphs[84], [("Hình 3: Bảng tổng hợp so sánh thực nghiệm 3 mẫu Agent qua 3 kịch bản từ benchmark.py", True, False, 10.5)], align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # -------------------------------------------------------------
    # SECTION 4: NHẬN XÉT VÀ PHÂN TÍCH CHUYÊN SÂU
    # -------------------------------------------------------------
    set_paragraph_content(doc.paragraphs[85], [("4) Nhận xét và phân tích chuyên sâu (In-depth Empirical Analysis):", True, False, 13.5)], space_after=8)

    set_paragraph_content(doc.paragraphs[86], [
        ("1. Hiệu suất số lượt suy luận (Turn Efficiency) & Chi phí ngữ cảnh:", True, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[87], [
        ("• Kịch bản thuận lợi (Happy Path): ", True, False, 12.0),
        ("Cả ReAct và Plan-then-Execute đều hoàn tất tối ưu trong đúng 4 lượt (Tìm chuyến bay → Khảo sát ghế → Tính hành lý → Đặt vé). Hybrid Agent tiêu tốn 7 lượt do chi phí phụ trội (overhead) để khởi tạo và đồng bộ trạng thái qua 3 Sub-Agent; bù lại mỗi lượt của Hybrid có kích thước prompt ngắn hơn đáng kể, tránh quá tải và phân tán chú ý của mô hình.", False, False, 12.0)
    ], space_after=6)

    set_paragraph_content(doc.paragraphs[88], [
        ("2. Khả năng thích ứng trước sự cố môi trường (Dynamic Fault Recovery):", True, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[89], [
        ("• Kịch bản cạn kiệt ghế (Seat Depletion): ", True, False, 12.0),
        ("Khi chuyến bay giá tốt nhất QH118 hết vé, tỷ lệ tự phục hồi của cả 3 mẫu đều đạt 100% (Pass). ReAct phản xạ nhanh nhất khi chỉ mất thêm 1 lượt thích ứng (tổng 5 lượt). Plan-then-Execute kích hoạt LLM Dynamic Re-planner để điều chỉnh 3 bước còn lại (tổng 5 lượt). Hybrid cô lập lỗi hoàn toàn trong Cột mốc 2 mà không làm xáo trộn Cột mốc 1 và 3 (tổng 7 lượt).", False, False, 12.0)
    ], space_after=6)

    set_paragraph_content(doc.paragraphs[90], [
        ("3. Vai trò sống còn của Harness Code đối với an toàn và độ tin cậy:", True, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[91], [
        ("• Bảo đảm an toàn tài chính (PermissionGuard): ", True, False, 12.0),
        ("Ở Kịch bản 3, 100% các mẫu Agent đều bị chặn đứng an toàn tại trạng thái WAIT_APPROVAL ngay trước lệnh book_ticket. Điều này chứng minh Agent không thể tự ý phát sinh chi phí nếu chưa có sự phê duyệt tường minh từ con người.", False, False, 12.0)
    ], space_after=3)
    set_paragraph_content(doc.paragraphs[92], [
        ("• Độ tin cậy nghiệm thu (CompletionSensor): ", True, False, 12.0),
        ("100% các lượt hoàn thành đều vượt qua kiểm tra đối chiếu trực tiếp 6 tiêu chí trong Mock DB. Cảm biến tính toán độc lập loại bỏ hoàn toàn khả năng mô hình tự ngụy tạo mã đặt chỗ ảo (Zero False-Positive Completions).", False, False, 12.0)
    ], space_after=8)

    # -------------------------------------------------------------
    # SECTION 5: KẾT LUẬN
    # -------------------------------------------------------------
    set_paragraph_content(doc.paragraphs[93], [("5) KẾT LUẬN (CONCLUSION)", True, False, 13.5)], space_after=8)

    set_paragraph_content(doc.paragraphs[94], [
        ("Quá trình thực nghiệm đã chứng minh rõ nét triết lý: ", False, False, 12.0),
        ("Xây dựng một hệ thống Agentic AI trong thực tế không chỉ dừng lại ở việc kết nối Prompt với LLM, mà nằm ở việc thiết kế khung Harness vững chắc bằng code bao bọc xung quanh.", True, False, 12.0)
    ], space_after=6)

    set_paragraph_content(doc.paragraphs[95], [
        ("Mô hình ngôn ngữ lớn (Model) đóng vai trò là \"bộ não\" suy luận linh hoạt, trong khi Harness đóng vai trò là \"bộ khung xương\" bảo đảm tính chính xác, an toàn tài chính và kiểm soát tài nguyên vận hành.", False, False, 12.0)
    ], space_after=6)

    set_paragraph_content(doc.paragraphs[96], [
        ("Khuyến nghị áp dụng: ", True, False, 12.0),
        ("ReAct phù hợp cho tác vụ nhanh, linh hoạt; Plan-then-Execute tối ưu cho các quy trình tuần tự cần kiểm duyệt kế hoạch; và Hybrid là giải pháp toàn diện cho các hệ sinh thái phức tạp, đa chặng đòi hỏi tính phân tầng và cô lập lỗi cao.", False, False, 12.0)
    ], space_after=12)

    doc.paragraphs[97].text = ""

    # Save modified document
    doc.save(docx_path)
    print("Document successfully rewritten and saved to:", docx_path)

if __name__ == "__main__":
    update_document()
